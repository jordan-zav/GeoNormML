"""Genera un PPTX vertical editable a partir de un JSON CPG XXIII.

La geometría se expresa en píxeles de un lienzo 1080 x 1920 y se convierte a
pulgadas a 144 dpi. Los contenidos pertenecen al JSON: el layout no contiene
datos de ningún trabajo.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

W, H, DPI = 1080, 1920, 144
BLUE = "003B70"; CYAN = "00A6CE"; INK = "15202B"; BG = "F5F8FA"; ORANGE = "D98324"; WHITE = "FFFFFF"; MUTED = "51606E"

def inch(px): return Inches(px / DPI)
def rgb(value): return RGBColor.from_string(value)

def shape(slide, x, y, w, h, fill, line=None, radius=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, inch(x), inch(y), inch(w), inch(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(line or fill)
    return s

def textbox(slide, text, x, y, w, h, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, margin=0, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = inch(margin)
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    for run in p.runs:
        run.font.name = "Arial"; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(color)
    return box

def image_or_placeholder(slide, path, caption, x, y, w, h):
    shape(slide, x, y, w, h, WHITE, CYAN, radius=True)
    source = Path(path) if path else None
    if source and source.exists():
        slide.shapes.add_picture(str(source), inch(x + 12), inch(y + 12), width=inch(w - 24), height=inch(h - 60))
    else:
        textbox(slide, "EVIDENCIA PENDIENTE\nNo se insertó contenido ficticio.", x + 28, y + 60, w - 56, h - 120, 13, ORANGE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    textbox(slide, caption or "Figura / captura real", x + 18, y + h - 42, w - 36, 26, 12, MUTED)

def header(slide, cfg, number):
    shape(slide, 0, 0, W, H, BG)
    shape(slide, 0, 0, W, 18, BLUE)
    textbox(slide, "XXIII CONGRESO PERUANO DE GEOLOGÍA", 54, 40, 610, 30, 12, BLUE, True)
    # Fixed, intentionally empty logo zones: source logos are supplied by the author.
    shape(slide, 790, 34, 105, 54, WHITE, CYAN, radius=True); textbox(slide, "LOGO\nUNI", 800, 41, 85, 40, 12, BLUE, True, PP_ALIGN.CENTER)
    shape(slide, 912, 34, 114, 54, WHITE, CYAN, radius=True); textbox(slide, "LOGO\nSGP/CPG", 920, 41, 98, 40, 12, BLUE, True, PP_ALIGN.CENTER)
    shape(slide, 54, 1810, 972, 2, CYAN)
    textbox(slide, cfg["short_title"], 54, 1830, 780, 28, 12, MUTED)
    textbox(slide, f"{number:02d}", 950, 1830, 76, 28, 12, BLUE, True, PP_ALIGN.RIGHT)

def bullets(slide, values, x, y, w, h):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h)); tf = box.text_frame; tf.clear(); tf.word_wrap = True
    for i, value in enumerate(values):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = value; p.level = 0; p.font.name = "Arial"; p.font.size = Pt(12); p.font.color.rgb = rgb(INK); p.space_after = Pt(8); p.bullet = True

def add_common_title(slide, data):
    textbox(slide, data.get("heading", "SECCIÓN"), 54, 130, 972, 60, 14, BLUE, True)
    body = data.get("body", "")
    if body: textbox(slide, body, 54, 215, 972, 150, 12, INK)

def add_slide(prs, cfg, data, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); header(slide, cfg, number); kind = data.get("kind", "context")
    if kind == "title":
        textbox(slide, cfg.get("title", "TÍTULO DEL TRABAJO"), 54, 240, 970, 190, 14, BLUE, True)
        textbox(slide, cfg.get("authors", "Autoría por confirmar"), 54, 465, 970, 60, 13, INK)
        textbox(slide, cfg.get("affiliation", "Afiliación institucional por confirmar"), 54, 540, 970, 60, 12, MUTED)
        image_or_placeholder(slide, data.get("visual"), data.get("caption", "Imagen científica representativa"), 54, 690, 972, 790)
        textbox(slide, data.get("body", "Resumen visual breve del trabajo."), 90, 1530, 900, 140, 13, INK, False, PP_ALIGN.CENTER)
        return
    add_common_title(slide, data)
    if kind == "method":
        steps = data.get("steps", [])
        for i, step in enumerate(steps[:4]):
            x = 54 + i * 243; shape(slide, x, 470, 210, 200, WHITE, CYAN, radius=True); textbox(slide, f"{i+1:02d}", x+18, 490, 174, 28, 12, ORANGE, True); textbox(slide, step, x+18, 535, 174, 106, 12, INK, True, PP_ALIGN.CENTER)
            if i < len(steps)-1: textbox(slide, "→", x+211, 545, 30, 30, 14, CYAN, True, PP_ALIGN.CENTER)
        image_or_placeholder(slide, data.get("visual"), data.get("caption", "Diagrama de arquitectura / flujo"), 54, 730, 972, 790)
    elif kind == "results":
        metrics = data.get("metrics", [])
        for i, metric in enumerate(metrics[:3]):
            x = 54 + i * 324; shape(slide, x, 410, 288, 150, BLUE, BLUE, radius=True); textbox(slide, metric.get("value", "—"), x+20, 432, 248, 44, 14, WHITE, True, PP_ALIGN.CENTER); textbox(slide, metric.get("label", "Métrica"), x+20, 492, 248, 30, 12, WHITE, False, PP_ALIGN.CENTER)
        visuals = data.get("visuals", [])[:2]
        for i, visual in enumerate(visuals): image_or_placeholder(slide, visual.get("path"), visual.get("caption"), 54 + i*486, 620, 450, 820)
        if data.get("bullets"): bullets(slide, data["bullets"], 54, 1485, 972, 250)
    elif kind == "conclusions":
        bullets(slide, data.get("bullets", []), 86, 400, 908, 550)
        image_or_placeholder(slide, data.get("visual"), data.get("caption", "Resultado o aplicación"), 54, 1040, 972, 600)
    else:
        if data.get("bullets"): bullets(slide, data["bullets"], 54, 400, 444, 320)
        image_or_placeholder(slide, data.get("visual"), data.get("caption", "Evidencia científica"), 540, 400, 486, 620)
        if data.get("body_secondary"): textbox(slide, data["body_secondary"], 54, 1120, 972, 250, 12, INK)

def main(config_path):
    config_path = Path(config_path); cfg = json.loads(config_path.read_text(encoding="utf-8"))
    if len(cfg.get("slides", [])) > 10: raise ValueError("El sistema admite un máximo de 10 pantallas por trabajo.")
    prs = Presentation(); prs.slide_width = inch(W); prs.slide_height = inch(H)
    for i, data in enumerate(cfg["slides"], start=1): add_slide(prs, cfg, data, i)
    target = config_path.with_suffix(".pptx"); prs.save(target); print(target)

if __name__ == "__main__":
    if len(sys.argv) != 2: raise SystemExit("Uso: python build_poster.py configs/trabajo.json")
    main(sys.argv[1])
