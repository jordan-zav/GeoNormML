"""Construye la primera version editorial del totem GeoNorm para CPG XXIII.

El lienzo usa 1080 x 1920 puntos para coincidir con los PDF de referencia de
MUSA y PetroPyQAPF. El resultado principal es un PPTX completamente editable.
"""

from __future__ import annotations

import os
from pathlib import Path
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
import qrcode


PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = Path(os.environ.get("GEONORMML_POSTER_OUTPUT_DIR", PROJECT_ROOT / "artifacts" / "posters" / "totem_v1"))
OUTPUT_PATH = OUTPUT_DIR / "Legacy_Totem_GeoNorm_CPG_XXIII.pptx"
ASSET_DIR = OUTPUT_DIR / "assets"

W, H, DPI = 1080, 1920, 72

PAPER = "F3EFE4"
PAPER_LIGHT = "FAF8F2"
INK = "17262D"
MUTED = "5E6B70"
NAVY = "102F3D"
TEAL = "287C8E"
TEAL_LIGHT = "DDECEF"
GREEN = "4C956C"
GREEN_LIGHT = "E3EFE7"
COPPER = "B66A4A"
COPPER_LIGHT = "F1E2DA"
GOLD = "D5A63B"
WHITE = "FFFFFF"
LINE = "D8D4CA"
RED_LIGHT = "F4E4E0"

MAP_PATH = PROJECT_ROOT / "docs" / "dataset_filtrado.jpg"
DASHBOARD_PATH = PROJECT_ROOT / "artifacts" / "dashboard_confiabilidad_actualizado.png"
CIPW_VOL_PATH = (
    PROJECT_ROOT
    / "artifacts"
    / "tas_comparison"
    / "concordancia_litologia_cipw_vol.png"
)
CIPW_PLU_PATH = (
    PROJECT_ROOT
    / "artifacts"
    / "tas_comparison"
    / "concordancia_litologia_cipw_plu.png"
)

APPROVED_TITLE = (
    "LEGACY CPG XXIII — "
    "GeoNormPy y GeoNormML: Un ecosistema abierto en Python e integración "
    "de algoritmos determinísticos para el cálculo normativo y clasificación "
    "de rocas ígneas en el Perú"
)


def u(px: float):
    return Inches(px / DPI)


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill, line=None, radius=0, width=1):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, u(x), u(y), u(w), u(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.color.rgb = color(line or fill)
    shape.line.width = Pt(width)
    if radius:
        try:
            shape.adjustments[0] = min(radius / min(w, h), 0.45)
        except (IndexError, ValueError):
            pass
    return shape


def line(slide, x, y, w, h, fill=LINE):
    return box(slide, x, y, w, h, fill, fill)


def text(
    slide,
    value,
    x,
    y,
    w,
    h,
    size=14,
    fill=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font="Arial",
    margin=0,
    italic=False,
):
    shape = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = u(margin)
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color(fill)
    return shape


def bullets(slide, items, x, y, w, line_height=76, size=15, accent=TEAL, text_fill=INK):
    for index, item in enumerate(items):
        yy = y + index * line_height
        box(slide, x, yy + 9, 12, 12, accent, accent, radius=4)
        text(slide, item, x + 28, yy, w - 28, line_height, size, text_fill)


def metric(slide, value, label, x, y, w=190, h=118, accent=TEAL, dark=False):
    fill = accent if dark else WHITE
    value_color = WHITE if dark else accent
    label_color = "E8F1F3" if dark else MUTED
    box(slide, x, y, w, h, fill, accent, radius=14, width=1.4)
    text(slide, value, x + 12, y + 16, w - 24, 43, 24, value_color, True, PP_ALIGN.CENTER)
    text(slide, label, x + 12, y + 69, w - 24, 28, 11, label_color, False, PP_ALIGN.CENTER)


def image_contain(slide, path: Path, x, y, w, h, pad=14, fill=WHITE, border=LINE):
    box(slide, x, y, w, h, fill, border, radius=12)
    if not path.exists():
        text(slide, f"EVIDENCIA PENDIENTE\n{path.name}", x, y, w, h, 14, COPPER, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        return
    with Image.open(path) as source:
        iw, ih = source.size
    scale = min((w - 2 * pad) / iw, (h - 2 * pad) / ih)
    dw, dh = iw * scale, ih * scale
    slide.shapes.add_picture(str(path), u(x + (w - dw) / 2), u(y + (h - dh) / 2), u(dw), u(dh))


def image_cover(slide, path: Path, x, y, w, h):
    box(slide, x, y, w, h, WHITE, WHITE, radius=12)
    if not path.exists():
        return
    with Image.open(path) as source:
        iw, ih = source.size
    target_ratio = w / h
    source_ratio = iw / ih
    pic = slide.shapes.add_picture(str(path), u(x), u(y), u(w), u(h))
    if source_ratio > target_ratio:
        visible = target_ratio / source_ratio
        pic.crop_left = pic.crop_right = (1 - visible) / 2
    else:
        visible = source_ratio / target_ratio
        pic.crop_top = pic.crop_bottom = (1 - visible) / 2


def base_slide(prs, number, section, accent=TEAL, accent_y=250, accent_h=480):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box(slide, 0, 0, W, H, PAPER, PAPER)
    box(slide, 0, 0, 72, H, NAVY, NAVY)
    box(slide, 0, accent_y, 72, accent_h, accent, accent)
    text(slide, "GEONORM", 112, 37, 250, 30, 12, NAVY, True)
    text(slide, f"PÁGINA {number:02d} / 10", 830, 37, 176, 30, 10, MUTED, False, PP_ALIGN.RIGHT)
    line(slide, 112, 83, 894, 2, LINE)
    text(slide, section.upper(), 112, 116, 894, 32, 11, accent, True)
    line(slide, 112, 1848, 894, 2, LINE)
    text(slide, "XXIII CONGRESO PERUANO DE GEOLOGÍA · 2026", 112, 1867, 520, 24, 9, MUTED)
    text(slide, "GeoNorm · J. Zavaleta", 760, 1867, 246, 24, 9, MUTED, False, PP_ALIGN.RIGHT)
    return slide


def title_block(slide, title_value, subtitle=None, y=172, title_h=150, size=28):
    text(slide, title_value, 112, y, 894, title_h, size, NAVY, True)
    if subtitle:
        text(slide, subtitle, 112, y + title_h + 12, 894, 78, 14, MUTED)


def card_title(slide, title_value, x, y, w, accent=TEAL):
    box(slide, x, y, w, 48, accent, accent)
    text(slide, title_value.upper(), x + 18, y + 13, w - 36, 24, 11, WHITE, True)


def slide_1(prs):
    slide = base_slide(
        prs,
        1,
        "Trabajo técnico científico · IA y Machine Learning",
        TEAL,
        260,
        600,
    )
    text(slide, APPROVED_TITLE, 112, 174, 894, 400, 31, NAVY, True)
    text(
        slide,
        "Geoquímica de roca total + norma CIPW + modelos supervisados bajo una cadena reproducible.",
        112,
        590,
        894,
        85,
        17,
        INK,
    )
    image_contain(slide, DASHBOARD_PATH, 112, 720, 894, 520, pad=10, fill=PAPER_LIGHT, border=TEAL)
    metric(slide, "2 887", "muestras · dataset v5", 112, 1285, 205, 120, TEAL, True)
    metric(slide, "19", "clases litológicas", 337, 1285, 205, 120, TEAL, True)
    metric(slide, "4", "modelos productivos", 562, 1285, 205, 120, TEAL, True)
    metric(slide, "16/16", "pruebas GeoNormPy", 787, 1285, 219, 120, TEAL, True)
    text(slide, "Jordan Zavaleta", 112, 1480, 620, 42, 18, NAVY, True)
    text(
        slide,
        "Universidad Nacional de Ingeniería · Facultad de Ingeniería Geológica, Minera y Metalúrgica\n"
        "Av. Túpac Amaru 210, Rímac, Lima, Perú · jordan.zavaleta.r@uni.pe",
        112,
        1532,
        760,
        88,
        12,
        MUTED,
    )
    box(slide, 112, 1680, 894, 112, NAVY, NAVY)
    text(
        slide,
        "PROPÓSITO",
        136,
        1700,
        150,
        24,
        10,
        GOLD,
        True,
    )
    text(
        slide,
        "Transformar composición en evidencia trazable sin reemplazar el criterio geológico.",
        300,
        1696,
        670,
        58,
        15,
        WHITE,
        True,
    )


def slide_2(prs):
    slide = base_slide(prs, 2, "Problema y objetivo", COPPER, 310, 500)
    title_block(
        slide,
        "Una cadena fragmentada no es reproducible",
        "El problema no era únicamente clasificar: era poder reconstruir cada decisión desde los óxidos de entrada hasta el modelo exportado.",
        title_h=92,
    )
    cards = [
        ("PLANILLAS", "Limpieza manual, versiones divergentes y reglas difíciles de auditar."),
        ("CIPW EXTERNO", "Dependencia de herramientas ajenas al pipeline y trazabilidad parcial."),
        ("MODELOS", "Selección por score aislado, sin separar calidad analítica y predictores."),
        ("TRANSFERENCIA", "Artefactos sin contrato explícito de variables, dominio ni advertencias."),
    ]
    for i, (heading, body) in enumerate(cards):
        yy = 470 + i * 210
        box(slide, 112, yy, 894, 166, WHITE, LINE, radius=10)
        box(slide, 112, yy, 190, 166, COPPER_LIGHT, COPPER_LIGHT)
        text(slide, f"0{i+1}", 134, yy + 27, 54, 35, 16, COPPER, True)
        text(slide, heading, 134, yy + 76, 145, 28, 12, NAVY, True)
        text(slide, body, 330, yy + 35, 640, 92, 15, INK)
    box(slide, 112, 1380, 894, 280, NAVY, NAVY, radius=12)
    text(slide, "OBJETIVO", 146, 1412, 180, 28, 11, GOLD, True)
    text(
        slide,
        "Construir un ecosistema abierto y auditable que integre curación, cálculo normativo CIPW, benchmarking y exportación de modelos por dominio litológico.",
        146,
        1462,
        820,
        106,
        19,
        WHITE,
        True,
    )
    text(
        slide,
        "Principio rector · QA/QC describe la calidad del dato; no se usa como atajo predictivo.",
        146,
        1582,
        820,
        40,
        12,
        "CDE0E6",
    )


def slide_3(prs):
    slide = base_slide(prs, 3, "Arquitectura reproducible", TEAL, 390, 780)
    title_block(
        slide,
        "De óxidos a modelos transferibles, sin saltos opacos",
        "Cinco etapas enlazadas conservan datos, criterios y artefactos de evaluación.",
        title_h=100,
    )
    stages = [
        ("01", "ENTRADA", "Óxidos mayores\nmetadatos\ncoordenadas", COPPER),
        ("02", "CURACIÓN", "nombres\nQA/QC\ndominios", GOLD),
        ("03", "GEONORMPY", "CIPW\nbalance\ndiagnósticos", TEAL),
        ("04", "GEONORMML", "features\nCV 5-fold\nbenchmark", GREEN),
        ("05", "SALIDA", "modelos\nmétricas\ncontrato", NAVY),
    ]
    for i, (number, heading, body, accent) in enumerate(stages):
        yy = 430 + i * 222
        box(slide, 160, yy, 798, 165, WHITE, accent, radius=12, width=1.5)
        box(slide, 160, yy, 125, 165, accent, accent)
        text(slide, number, 176, yy + 50, 92, 46, 22, WHITE, True, PP_ALIGN.CENTER)
        text(slide, heading, 320, yy + 30, 235, 28, 12, accent, True)
        text(slide, body, 320, yy + 72, 280, 72, 14, INK)
        if i < 4:
            text(slide, "↓", 500, yy + 170, 80, 48, 22, accent, True, PP_ALIGN.CENTER)
        if i in (2, 3):
            note = "motor determinístico" if i == 2 else "selección probabilística"
            box(slide, 664, yy + 48, 250, 58, PAPER_LIGHT, accent, radius=8)
            text(slide, note, 680, yy + 65, 218, 24, 11, accent, True, PP_ALIGN.CENTER)
    box(slide, 112, 1595, 894, 112, TEAL_LIGHT, TEAL_LIGHT, radius=10)
    text(
        slide,
        "Resultado: cada predicción conserva dominio, conjunto de variables, algoritmo, métricas y columnas requeridas.",
        142,
        1625,
        834,
        50,
        14,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )


def slide_4(prs):
    slide = base_slide(prs, 4, "Datos y trazabilidad", TEAL, 560, 760)
    title_block(
        slide,
        "Dataset v5: 2 887 muestras y 19 clases",
        "El corte espacial actualizado amplía el conjunto aceptado originalmente y conserva la separación operativa entre rocas volcánicas y plutónicas.",
        title_h=92,
    )
    image_contain(slide, MAP_PATH, 112, 425, 540, 1100, pad=12, fill=WHITE, border=TEAL)
    metric(slide, "1 821", "volcánicas · 9 clases", 690, 470, 282, 130, TEAL, True)
    metric(slide, "1 066", "plutónicas · 10 clases", 690, 625, 282, 130, COPPER, True)
    metric(slide, "100 %", "CIPW calculado", 690, 780, 282, 130, GREEN, True)
    box(slide, 690, 960, 282, 340, WHITE, LINE, radius=10)
    card_title(slide, "Criterios", 690, 960, 282, NAVY)
    bullets(
        slide,
        [
            "Nombres litológicos estandarizados.",
            "Clases raras consolidadas con criterio petrológico.",
            "QA/QC separado del entrenamiento.",
            "Identificador final estable por muestra.",
        ],
        716,
        1035,
        230,
        line_height=63,
        size=12,
        accent=TEAL,
    )
    box(slide, 690, 1335, 282, 190, COPPER_LIGHT, COPPER_LIGHT, radius=10)
    text(slide, "FUENTES", 714, 1365, 230, 25, 11, COPPER, True)
    text(slide, "INGEMMET · GEOROC · USGS", 714, 1410, 230, 42, 14, NAVY, True)
    text(slide, "Cobertura regional del Perú", 714, 1470, 230, 28, 11, MUTED)
    box(slide, 112, 1570, 860, 120, NAVY, NAVY, radius=10)
    text(slide, "TRAZABILIDAD", 140, 1598, 160, 26, 10, GOLD, True)
    text(slide, "dataset → split → benchmark → modelo → catálogo", 316, 1594, 620, 38, 16, WHITE, True)
    text(slide, "La versión y el dominio acompañan cada artefacto exportado.", 316, 1640, 620, 28, 11, "D5E5E9")


def slide_5(prs):
    slide = base_slide(prs, 5, "GeoNormPy · motor CIPW", GREEN, 360, 790)
    title_block(
        slide,
        "CIPW determinístico, diagnosticable y reproducible",
        "GeoNormPy reemplaza la etapa externa y devuelve minerales normativos junto con controles interpretativos.",
        title_h=100,
    )
    steps = [
        ("01", "Normalización anhidra", "Reescala los óxidos disponibles."),
        ("02", "Especiación del hierro", "Maneja FeO, Fe₂O₃ y FeOT."),
        ("03", "Asignación normativa", "Distribuye fases en secuencia CIPW."),
        ("04", "Saturación", "Reporta sílice y estado de alúmina."),
        ("05", "Balance de masa", "Cuantifica el error de reconstrucción."),
    ]
    for i, (number, heading, body) in enumerate(steps):
        yy = 430 + i * 170
        box(slide, 112, yy, 630, 132, WHITE, LINE, radius=10)
        box(slide, 112, yy, 92, 132, GREEN, GREEN)
        text(slide, number, 128, yy + 44, 60, 34, 17, WHITE, True, PP_ALIGN.CENTER)
        text(slide, heading, 230, yy + 26, 470, 28, 13, NAVY, True)
        text(slide, body, 230, yy + 67, 470, 42, 12, MUTED)
    metric(slide, "16/16", "pruebas aprobadas", 778, 455, 208, 126, GREEN, True)
    metric(slide, "99.45 %", "balance ≤ 0.10", 778, 610, 208, 126, GREEN, True)
    metric(slide, "2 887", "cálculos sin error", 778, 765, 208, 126, GREEN, True)
    box(slide, 778, 940, 208, 342, GREEN_LIGHT, GREEN, radius=10)
    text(slide, "SALIDAS", 800, 970, 164, 26, 11, GREEN, True)
    bullets(
        slide,
        ["Minerales normativos", "Saturación en sílice", "Estado de alúmina", "Modo de hierro", "Advertencias", "Error de balance"],
        802,
        1020,
        160,
        line_height=46,
        size=11,
        accent=GREEN,
    )
    box(slide, 112, 1340, 874, 310, NAVY, NAVY, radius=12)
    text(slide, "ALCANCE CORRECTO", 144, 1374, 250, 26, 11, GOLD, True)
    text(
        slide,
        "La norma CIPW es una mineralogía ideal calculada desde química de roca total. Sirve como contraste composicional y descriptor reproducible; no sustituye una medición modal ni una clasificación QAPF formal.",
        144,
        1424,
        810,
        130,
        17,
        WHITE,
        True,
    )
    text(slide, "Umbral operativo de balance: error ≤ 0.10.", 144, 1580, 810, 30, 11, "D5E5E9")


def slide_6(prs):
    slide = base_slide(prs, 6, "GeoNormML · diseño experimental", TEAL, 410, 830)
    title_block(
        slide,
        "Comparar modelos sin confundir contexto con fuga de información",
        "La selección combina desempeño, disponibilidad de variables, estabilidad y coherencia petrológica.",
        title_h=100,
    )
    sets = [
        ("OXIDES_CORE", "Óxidos mayores, LOI y suma de óxidos", COPPER),
        ("CIPW", "Minerales normativos calculados por GeoNormPy", GREEN),
        ("COMBINED_CORE", "Óxidos + CIPW · baseline transferible", TEAL),
        ("ENRICHED_SPATIAL", "Core + índices + latitud + distancia a la fosa", NAVY),
    ]
    for i, (heading, body, accent) in enumerate(sets):
        yy = 430 + i * 174
        box(slide, 112, yy, 894, 132, WHITE, accent, radius=9, width=1.4)
        box(slide, 112, yy, 218, 132, accent, accent)
        text(slide, heading, 128, yy + 38, 186, 48, 11, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(slide, body, 358, yy + 40, 612, 54, 14, INK, True)
    card_title(slide, "Validación", 112, 1160, 278, TEAL)
    card_title(slide, "Algoritmos", 420, 1160, 278, GREEN)
    card_title(slide, "Variables excluidas", 728, 1160, 278, COPPER)
    box(slide, 112, 1208, 278, 320, WHITE, LINE, radius=8)
    box(slide, 420, 1208, 278, 320, WHITE, LINE, radius=8)
    box(slide, 728, 1208, 278, 320, WHITE, LINE, radius=8)
    bullets(slide, ["CV estratificada", "5 pliegues", "Métricas multiclase", "Semilla fija"], 136, 1250, 230, 58, 12, TEAL)
    bullets(slide, ["Random Forest", "XGBoost", "SVM + PCA", "Imputación por mediana"], 444, 1250, 230, 58, 12, GREEN)
    bullets(slide, ["Quality_Flag", "Geol_Unit", "Metal_Belt", "UTM_Zone"], 752, 1250, 230, 58, 12, COPPER)
    box(slide, 112, 1570, 894, 118, COPPER_LIGHT, COPPER_LIGHT, radius=10)
    text(slide, "REGLA", 140, 1600, 100, 26, 10, COPPER, True)
    text(
        slide,
        "Una variable no se acepta solo porque eleva la exactitud: debe ser disponible, auditable y geológicamente defendible.",
        250,
        1594,
        720,
        54,
        14,
        NAVY,
        True,
    )


def comparison_bar(slide, label, baseline, spatial, x, y, w, accent):
    text(slide, label, x, y, w, 26, 11, NAVY, True)
    bar_x, bar_y, bar_w = x, y + 38, w
    box(slide, bar_x, bar_y, bar_w, 26, LINE, LINE, radius=5)
    box(slide, bar_x, bar_y, bar_w * baseline, 26, accent, accent, radius=5)
    text(slide, f"{baseline:.3f}", bar_x + bar_w * baseline - 62, bar_y + 3, 58, 18, 9, WHITE, True, PP_ALIGN.RIGHT)
    box(slide, bar_x, bar_y + 42, bar_w, 26, LINE, LINE, radius=5)
    box(slide, bar_x, bar_y + 42, bar_w * spatial, 26, GREEN, GREEN, radius=5)
    text(slide, f"{spatial:.3f}", bar_x + bar_w * spatial - 62, bar_y + 45, 58, 18, 9, WHITE, True, PP_ALIGN.RIGHT)


def slide_7(prs):
    slide = base_slide(prs, 7, "Resultados · baselines", TEAL, 420, 760)
    title_block(
        slide,
        "Sin coordenadas: dos baselines transferibles",
        "Los modelos core constituyen la referencia principal cuando no existe contexto espacial válido.",
        title_h=95,
    )
    cards = [
        (112, "VOLCÁNICAS", "XGBoost", "combined_core", TEAL, ["0.806", "0.956", "0.718", "0.803"]),
        (574, "PLUTÓNICAS", "Random Forest", "combined_core", COPPER, ["0.615", "0.870", "0.523", "0.586"]),
    ]
    for x, domain, model, features, accent, values in cards:
        box(slide, x, 430, 432, 560, WHITE, accent, radius=12, width=1.4)
        box(slide, x, 430, 432, 76, accent, accent)
        text(slide, domain, x + 24, 453, 384, 28, 14, WHITE, True)
        text(slide, model, x + 28, 545, 376, 46, 24, NAVY, True)
        text(slide, features, x + 28, 605, 376, 30, 12, accent, True)
        labels = ["CA", "AUC macro", "MCC", "F1 pond."]
        for i, (value, label) in enumerate(zip(values, labels)):
            xx = x + 28 + (i % 2) * 188
            yy = 680 + (i // 2) * 122
            metric(slide, value, label, xx, yy, 164, 98, accent, False)
        finding = "Mayor separabilidad geoquímica." if domain == "VOLCÁNICAS" else "Mayor continuidad composicional."
        box(slide, x + 28, 918, 376, 48, PAPER_LIGHT, accent, radius=6)
        text(slide, finding, x + 40, 933, 352, 22, 11, accent, True, PP_ALIGN.CENTER)
    box(slide, 112, 1050, 894, 440, PAPER_LIGHT, LINE, radius=10)
    box(slide, 140, 1086, 42, 16, TEAL, TEAL, radius=3)
    text(slide, "BASELINE", 196, 1082, 126, 24, 10, MUTED, True)
    box(slide, 350, 1086, 42, 16, GREEN, GREEN, radius=3)
    text(slide, "CON CONTEXTO ESPACIAL", 406, 1082, 240, 24, 10, MUTED, True)
    comparison_bar(slide, "Exactitud · VOL", 0.8056, 0.8281, 140, 1145, 810, TEAL)
    comparison_bar(slide, "MCC · VOL", 0.7183, 0.7506, 140, 1255, 810, TEAL)
    comparison_bar(slide, "Exactitud · PLU", 0.6154, 0.6567, 140, 1365, 810, COPPER)
    box(slide, 112, 1540, 894, 150, NAVY, NAVY, radius=10)
    text(slide, "LECTURA", 142, 1570, 110, 25, 10, GOLD, True)
    text(
        slide,
        "La menor exactitud plutónica es coherente con fronteras que dependen de proporciones modales y textura, no solo de química total.",
        270,
        1565,
        700,
        72,
        14,
        WHITE,
        True,
    )


def slide_8(prs):
    slide = base_slide(prs, 8, "Resultados · contexto espacial", GREEN, 500, 820)
    title_block(
        slide,
        "El contexto espacial mejora el desempeño dentro del dominio muestreado",
        "Latitud y distancia a la fosa aportan estructura regional sin incorporar unidad geológica ni cinturón metalogenético.",
        title_h=128,
    )
    image_contain(slide, DASHBOARD_PATH, 112, 455, 894, 620, pad=8, fill=WHITE, border=GREEN)
    metric(slide, "+2.25 pp", "exactitud VOL", 112, 1120, 266, 122, GREEN, True)
    metric(slide, "+4.13 pp", "exactitud PLU", 406, 1120, 266, 122, GREEN, True)
    metric(slide, "0.751", "MCC espacial VOL", 700, 1120, 306, 122, GREEN, True)
    box(slide, 112, 1290, 430, 295, GREEN_LIGHT, GREEN, radius=10)
    card_title(slide, "Qué aporta", 112, 1290, 430, GREEN)
    bullets(
        slide,
        ["Posición longitudinal del arco.", "Distancia transversal a la fosa.", "Índices petrológicos derivados."],
        142,
        1365,
        370,
        line_height=64,
        size=13,
        accent=GREEN,
    )
    box(slide, 576, 1290, 430, 295, RED_LIGHT, COPPER, radius=10)
    card_title(slide, "Límite", 576, 1290, 430, COPPER)
    bullets(
        slide,
        ["CV estratificada, no espacial.", "Puede captar autocorrelación regional.", "Requiere validación por bloques."],
        606,
        1365,
        370,
        line_height=64,
        size=13,
        accent=COPPER,
    )
    box(slide, 112, 1620, 894, 86, NAVY, NAVY, radius=8)
    text(
        slide,
        "Uso recomendado: muestras con coordenadas válidas y contexto compatible con los Andes peruanos.",
        140,
        1643,
        838,
        36,
        13,
        WHITE,
        True,
        PP_ALIGN.CENTER,
    )


def slide_9(prs):
    slide = base_slide(prs, 9, "Discusión y límites", COPPER, 520, 780)
    title_block(
        slide,
        "CIPW describe composición; la litología conserva su autoridad geológica",
        "Las matrices muestran saturación normativa por clase reportada. No son matrices de validación modal ni sustituyen QAPF.",
        title_h=126,
    )
    image_contain(slide, CIPW_VOL_PATH, 112, 460, 894, 475, pad=8, fill=WHITE, border=TEAL)
    image_contain(slide, CIPW_PLU_PATH, 112, 980, 894, 475, pad=8, fill=WHITE, border=COPPER)
    box(slide, 112, 1500, 894, 200, NAVY, NAVY, radius=10)
    text(slide, "TRES LÍMITES EXPLÍCITOS", 140, 1530, 320, 25, 10, GOLD, True)
    bullets(
        slide,
        [
            "La litología es reportada; no constituye una nueva medición modal.",
            "Las probabilidades ML apoyan la revisión; no asignan autoridad petrográfica.",
            "La generalización espacial externa aún debe demostrarse.",
        ],
        142,
        1572,
        820,
        line_height=42,
        size=11,
        accent=GOLD,
        text_fill=WHITE,
    )


def slide_10(prs, qr_path):
    slide = base_slide(prs, 10, "Conclusiones y acceso", TEAL, 390, 900)
    title_block(
        slide,
        "GeoNorm convierte una cadena dispersa en evidencia auditable",
        "El aporte central es metodológico: separar cálculo determinístico, aprendizaje probabilístico y decisión geológica.",
        title_h=116,
    )
    conclusions = [
        ("01", "GeoNormPy integra CIPW, diagnósticos y balance de masa en Python."),
        ("02", "GeoNormML selecciona modelos por dominio y conserva sus contratos de entrada."),
        ("03", "Los baselines operan sin coordenadas; los modelos espaciales añaden contexto regional."),
        ("04", "La suite apoya —no reemplaza— el criterio petrológico y la autoridad modal QAPF."),
    ]
    for i, (number, body) in enumerate(conclusions):
        yy = 430 + i * 180
        box(slide, 112, yy, 894, 136, WHITE, TEAL, radius=10)
        box(slide, 112, yy, 100, 136, TEAL, TEAL)
        text(slide, number, 130, yy + 47, 64, 36, 18, WHITE, True, PP_ALIGN.CENTER)
        text(slide, body, 246, yy + 34, 720, 70, 15, NAVY, True, valign=MSO_ANCHOR.MIDDLE)
    box(slide, 112, 1195, 560, 410, NAVY, NAVY, radius=12)
    text(slide, "CÓDIGO ABIERTO", 146, 1230, 260, 26, 11, GOLD, True)
    text(slide, "GeoNormPy", 146, 1276, 400, 50, 27, WHITE, True)
    text(slide, "pip install geonormpy", 146, 1342, 460, 40, 17, "BFE3E6", True)
    text(slide, "github.com/jordan-zav/GeoNormPy", 146, 1406, 460, 34, 12, WHITE, True)
    text(
        slide,
        "Motor CIPW disponible bajo licencia abierta. GeoNormML se transfiere mediante artefactos versionados para uso académico.",
        146,
        1472,
        470,
        84,
        12,
        "D4E4E8",
    )
    image_contain(slide, qr_path, 710, 1195, 296, 296, pad=28, fill=WHITE, border=TEAL)
    text(slide, "Repositorio GeoNormPy", 710, 1510, 296, 28, 11, TEAL, True, PP_ALIGN.CENTER)
    text(slide, "Contacto", 710, 1580, 296, 24, 10, MUTED, True, PP_ALIGN.CENTER)
    text(slide, "jordan.zavaleta.r@uni.pe", 690, 1615, 336, 30, 11, NAVY, True, PP_ALIGN.CENTER)
    box(slide, 112, 1680, 894, 72, TEAL_LIGHT, TEAL_LIGHT, radius=8)
    text(
        slide,
        "Siguiente etapa científica · validación externa y evaluación por bloques espaciales.",
        136,
        1701,
        846,
        30,
        12,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )


def build_qr() -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / "qr_geonormpy.png"
    qr = qrcode.QRCode(version=5, box_size=10, border=2)
    qr.add_data("https://github.com/jordan-zav/GeoNormPy")
    qr.make(fit=True)
    image = qr.make_image(fill_color="#102F3D", back_color="white")
    image.save(path)
    return path


def main() -> Path:
    if os.environ.get("GEONORMML_ALLOW_LEGACY_POSTER") != "1":
        raise RuntimeError("Historical CPG generator: set GEONORMML_ALLOW_LEGACY_POSTER=1 explicitly. Use current JSON configs for current results.")

    required = [MAP_PATH, DASHBOARD_PATH, CIPW_VOL_PATH, CIPW_PLU_PATH]
    missing = [path for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Faltan recursos: " + ", ".join(str(path) for path in missing))

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    qr_path = build_qr()
    prs = Presentation()
    prs.slide_width = u(W)
    prs.slide_height = u(H)
    slide_1(prs)
    slide_2(prs)
    slide_3(prs)
    slide_4(prs)
    slide_5(prs)
    slide_6(prs)
    slide_7(prs)
    slide_8(prs)
    slide_9(prs)
    slide_10(prs, qr_path)
    prs.save(OUTPUT_PATH)
    print(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        raise
